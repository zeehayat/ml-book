"""
Build a "concept + exercise" PowerPoint presentation for Chapter 3
(The Anatomy of a Tensor & Compute Hardware).

Each content slide covers exactly one concept, then poses one exercise to
practice it immediately. Speaker notes hold the exercise's worked answer.
The final slide is a comprehensive mastery checklist spanning the chapter.

Run:
    python3 build_chapter3_practice_presentation.py

Produces:
    Chapter3_Tensors_Practice_Presentation.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------------------
# Palette (matches the Chapter 1/2 practice decks for a consistent series look)
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1B, 0x26, 0x4F)
AMBER = RGBColor(0xB8, 0x6A, 0x00)
AMBER_BG = RGBColor(0xFD, 0xF2, 0xE2)
LIGHT_BG = RGBColor(0xF7, 0xF9, 0xFC)
TEXT_DARK = RGBColor(0x22, 0x22, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MASTERY_BG = RGBColor(0xEA, 0xF6, 0xF4)
TEAL = RGBColor(0x0D, 0x94, 0x88)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# ---------------------------------------------------------------------------
# Content: one concept + one exercise per slide
# ---------------------------------------------------------------------------
SLIDES = [
    {
        "title": "Chapter 3: The Anatomy of a Tensor & Compute Hardware",
        "subtitle": "One concept, one exercise per slide — practice as you go",
        "concept": [],
        "exercise": None,
        "notes": "Welcome slide. Every tensor library (NumPy, PyTorch, JAX) represents an array as a flat memory buffer plus a small set of metadata. This chapter builds that model from first principles.",
    },
    {
        "title": "Concept: A Tensor Is a Flat Buffer Plus Metadata",
        "concept": [
            "The actual numbers live in one flat, one-dimensional buffer in memory",
            "Four metadata fields describe how to read that buffer as an N-D array: shape, strides, offset, dtype",
            "The \"shape\" you picture (a 3×4 grid) is a logical view, not the physical layout",
            "Two tensors can share the exact same buffer while looking completely different",
        ],
        "exercise": "A 2×3 matrix is stored as the flat buffer [10, 20, 30, 40, 50, 60]. If this is C-contiguous (row-major), what are the two rows of the logical matrix?",
        "notes": "Answer: row 0 = [10, 20, 30], row 1 = [40, 50, 60] — the first 3 elements of the flat buffer form row 0, the next 3 form row 1, read left to right, top to bottom.",
    },
    {
        "title": "Concept: Shape, Dimension, and Rank",
        "concept": [
            "Shape: a tuple stating the size along each axis, e.g. (3, 4) means 3 rows, 4 columns",
            "Rank (or ndim): the number of axes — a matrix has rank 2, a vector has rank 1",
            "The total element count is the product of all shape entries",
        ],
        "exercise": "A tensor has shape (2, 3, 4). What is its rank, and how many total elements does it hold?",
        "notes": "Answer: rank = 3 (three axes). Total elements = 2 × 3 × 4 = 24.",
    },
    {
        "title": "Concept: Strides — The Distance Between Elements",
        "concept": [
            "A stride is the number of buffer positions to skip to move one step along an axis",
            "strides = (s0, s1, ...) — one stride value per axis",
            "A stride of 1 means \"adjacent in memory\"; a large stride means \"far apart\"",
            "Strides — not shape — are what actually determine memory access cost",
        ],
        "exercise": "For shape (3, 4) with strides (4, 1), how many buffer positions do you skip to move from element [1,2] to element [2,2] (one step along axis 0)?",
        "notes": "Answer: 4 positions — moving one step along axis 0 (the row axis) skips exactly stride[0] = 4 buffer positions, since each row has 4 elements.",
    },
    {
        "title": "Concept: C-Contiguous (Row-Major) Strides",
        "concept": [
            "Row-major: elements of the LAST axis are stored consecutively in memory",
            "Stride formula: stride[k] = product of all shape entries AFTER axis k",
            "This is NumPy's and PyTorch's default layout",
        ],
        "exercise": "Derive the C-contiguous strides for a tensor with shape (5, 6, 7).",
        "notes": "Answer: stride[2] = 1 (last axis, always 1 element apart). stride[1] = shape[2] = 7. stride[0] = shape[1]*shape[2] = 6*7 = 42. Strides = (42, 7, 1).",
    },
    {
        "title": "Concept: F-Contiguous (Column-Major) Strides",
        "concept": [
            "Column-major: elements of the FIRST axis are stored consecutively in memory",
            "Stride formula: stride[k] = product of all shape entries BEFORE axis k",
            "Fortran, MATLAB, and some linear algebra libraries default to this layout",
        ],
        "exercise": "Derive the F-contiguous strides for the same shape (5, 6, 7), and state which single stride value is guaranteed to be 1 in each layout (C vs F).",
        "notes": "Answer: stride[0]=1, stride[1]=shape[0]=5, stride[2]=shape[0]*shape[1]=30. Strides = (1, 5, 30). In C-contiguous layout, the LAST axis's stride is always 1; in F-contiguous layout, the FIRST axis's stride is always 1.",
    },
    {
        "title": "Concept: The Flat Index Formula",
        "concept": [
            "Any logical index (i0, i1, ..., ik) maps to one flat buffer position",
            "flat_index = offset + i0*stride0 + i1*stride1 + ... + ik*stride_k",
            "The offset lets a view start partway into a shared buffer",
        ],
        "exercise": "A tensor has offset=2, shape=(3,4), strides=(4,1). What is the flat buffer index of logical element [1, 2]?",
        "notes": "Answer: flat_index = 2 + 1*4 + 2*1 = 2 + 4 + 2 = 8.",
    },
    {
        "title": "Concept: Slicing Is a Zero-Copy View",
        "concept": [
            "Slicing a tensor does NOT copy the underlying buffer",
            "It only computes new shape, strides, and offset that point into the SAME buffer",
            "Mutating a slice therefore mutates the original tensor too",
        ],
        "exercise": "You slice a large tensor with x[10:20], then modify one element of the result. Does the original tensor x change? Why or why not?",
        "notes": "Answer: yes, it changes — the slice is a view sharing the same underlying buffer as x, with adjusted offset/shape/strides. Writing to the view writes to the same memory the original tensor reads from.",
    },
    {
        "title": "Concept: Transpose Is a Zero-Copy Stride Permutation",
        "concept": [
            "Transposing a tensor just swaps its strides — no data moves at all",
            "A (3,4) C-contiguous matrix with strides (4,1), transposed, becomes shape (4,3) with strides (1,4)",
            "The transposed result is no longer C-contiguous — it's now F-contiguous relative to the original layout",
        ],
        "exercise": "A matrix has shape (3,4) and strides (4,1). After transposing, what are the new shape and strides? Is the result C-contiguous?",
        "notes": "Answer: new shape = (4,3), new strides = (1,4) — shape and strides are simply reversed. It is NOT C-contiguous (the last axis's stride is 4, not 1) — this is exactly why some operations after a transpose silently trigger a copy.",
    },
    {
        "title": "Concept: Reshape — Sometimes a View, Sometimes a Copy",
        "concept": [
            "reshape() returns a view whenever the existing strides make the new shape reachable without moving data",
            "If not possible (e.g. reshaping a transposed, non-contiguous tensor), it silently copies",
            "view() (PyTorch) refuses to reshape when a copy would be needed, forcing you to notice",
        ],
        "exercise": "Why does calling .view() on a transposed PyTorch tensor often raise a RuntimeError, while .reshape() on the same tensor works without complaint?",
        "notes": "Answer: a transposed tensor is usually non-contiguous, so its existing strides cannot represent the new shape as a pure view — .view() refuses and raises an error to make this explicit, while .reshape() silently falls back to copying the data into a new contiguous buffer first.",
    },
    {
        "title": "Concept: Broadcasting",
        "concept": [
            "Broadcasting lets operations combine tensors of different (but compatible) shapes",
            "Two dimensions are compatible if they are equal, or one of them is 1",
            "A size-1 dimension is \"stretched\" conceptually — using a stride of 0, not a real copy",
        ],
        "exercise": "Can a tensor of shape (5, 1) be broadcast against a tensor of shape (5, 3)? What about shape (4, 3) against (5, 3)?",
        "notes": "Answer: (5,1) and (5,3) ARE compatible — the size-1 dimension broadcasts to 3. (4,3) and (5,3) are NOT compatible — 4 and 5 are unequal and neither is 1.",
    },
    {
        "title": "Concept: Cache Locality and the AMAT Model",
        "concept": [
            "CPUs fetch memory in fixed-size cache lines (typically 64 bytes), not single values",
            "Sequential (stride-1) access reuses a fetched cache line many times — cheap",
            "Large-stride access fetches a new cache line almost every time — expensive",
            "This is a simplified model — real hardware has prefetching and multiple cache levels",
        ],
        "exercise": "For float32 values (4 bytes) and a 64-byte cache line, how many elements fit in one cache line? If you traverse a large matrix with a stride of 1000 elements, roughly what fraction of that cache line's capacity do you actually use per fetch?",
        "notes": "Answer: 64/4 = 16 elements fit per cache line. With a stride of 1000 (far larger than 16), essentially every access lands in a new cache line — you use only 1 out of 16 possible elements per fetch, about 6.25% utilization, under this simplified model.",
    },
    {
        "title": "Concept: dtype and Memory Footprint",
        "concept": [
            "Every element's byte width depends on its dtype: float32=4 bytes, float64=8 bytes, int8=1 byte",
            "Total buffer size = (number of elements) × (bytes per element) — strides do not change this",
            "Downcasting (e.g. float64 → float16) saves memory but can silently lose precision",
        ],
        "exercise": "A tensor has shape (1000, 1000) and dtype float32. How many megabytes does its buffer occupy? How much would switching to float16 save?",
        "notes": "Answer: 1000*1000 = 1,000,000 elements * 4 bytes = 4,000,000 bytes ≈ 3.81 MB. At float16 (2 bytes), it would be about 1.91 MB — exactly half, since only the byte width changed.",
    },
    {
        "title": "Concept: Common Mistake — Assuming a Slice Is an Independent Copy",
        "concept": [
            "Beginners often slice a tensor, modify the slice, and are surprised the original changed",
            "This is not a bug — it's the zero-copy view behavior from earlier in this deck",
            "Fix: call .copy() (NumPy) or .clone() (PyTorch) explicitly when independence is required",
        ],
        "exercise": "You want to extract a sub-matrix, modify it freely, and guarantee the original tensor is untouched. What single method call fixes this, and where must you place it?",
        "notes": "Answer: call .copy() (or .clone() in PyTorch) on the slice immediately after slicing, before any modification — e.g. sub = x[1:3].copy(). Without it, sub remains a view into x's buffer.",
    },
    {
        "title": "Concept: Common Mistake — Off-by-One Stride Errors",
        "concept": [
            "Manually computing strides for a shape is a common source of silent bugs",
            "A single wrong stride value produces a tensor that reads plausible-looking but WRONG data",
            "Always verify manual stride math against a known library's actual .strides output",
        ],
        "exercise": "You compute C-contiguous strides for shape (3, 4) by hand and get (3, 1). Is this correct? If not, find the error.",
        "notes": "Answer: incorrect. stride[1] should be 1 (correct), but stride[0] should be the product of shape entries AFTER axis 0, i.e. shape[1] = 4, not 3 (the size of axis 0 itself). The correct strides are (4, 1) — a classic off-by-one confusion between \"this axis's own size\" and \"the size of everything after it.\"",
    },
    {
        "title": "Mastery Check: Are You Ready to Move On?",
        "concept": [
            "If you can answer every question below confidently, you have mastered Chapter 3.",
        ],
        "exercise": None,
        "mastery_questions": [
            "Explain the four pieces of metadata (shape, strides, offset, dtype) that turn a flat buffer into a tensor.",
            "Derive the C-contiguous strides for a tensor of shape (D, H, W) from scratch.",
            "Derive the F-contiguous strides for the same shape and explain the difference from C-contiguous.",
            "Write the flat index formula and use it to compute a specific element's buffer position by hand.",
            "Explain why slicing is zero-copy and why mutating a slice can mutate the original tensor.",
            "Explain why transposing a tensor changes its strides but touches no data.",
            "Explain why .reshape() sometimes silently copies data while .view() refuses to.",
            "State the broadcasting compatibility rule and apply it to two given shapes.",
            "Explain, using the AMAT model, why a stride-1 traversal is faster than a large-stride traversal.",
            "Compute the memory footprint of a tensor given its shape and dtype.",
            "Diagnose a bug where a modified 'copy' of a tensor unexpectedly altered the original.",
            "Find and correct an off-by-one error in a hand-computed stride formula.",
            "Implement a minimal pure-Python flat-array class supporting shape, strides, and indexing.",
        ],
        "notes": "This is a self-check slide. If any question causes hesitation, return to that concept's slide and its exercise before moving to Chapter 4.",
    },
]


def add_slide_number(slide, n, total):
    box = slide.shapes.add_textbox(SLIDE_W - Inches(1.3), SLIDE_H - Inches(0.5), Inches(1.0), Inches(0.4))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{n} / {total}"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def set_notes(slide, text):
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = text


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]
    total = len(SLIDES)

    for idx, s in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG

        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.15))
        header.fill.solid()
        header.fill.fore_color.rgb = NAVY
        header.line.fill.background()
        header.shadow.inherit = False
        title_box = header.text_frame
        title_box.word_wrap = True
        title_box.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = title_box.paragraphs[0]
        p.text = s["title"]
        p.font.size = Pt(28) if idx > 1 else Pt(32)
        p.font.bold = True
        p.font.color.rgb = WHITE
        title_box.margin_left = Inches(0.5)

        if idx == 1:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), SLIDE_W - Inches(1.6), Inches(1.5))
            tf = sub_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = s["subtitle"]
            run.font.size = Pt(22)
            run.font.color.rgb = TEAL
            run.font.italic = True
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.7), Inches(3.0), Pt(4))
            accent.fill.solid()
            accent.fill.fore_color.rgb = TEAL
            accent.line.fill.background()

        elif "mastery_questions" in s:
            body_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.35), SLIDE_W - Inches(1.2), Inches(0.5))
            tf = body_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = s["concept"][0]
            run.font.size = Pt(15)
            run.font.italic = True
            run.font.color.rgb = TEAL

            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.85), SLIDE_W - Inches(1.2), Inches(5.35)
            )
            box.fill.solid()
            box.fill.fore_color.rgb = MASTERY_BG
            box.line.color.rgb = TEAL
            box.line.width = Pt(1.25)
            box.shadow.inherit = False
            ctf = box.text_frame
            ctf.word_wrap = True
            ctf.margin_left = Inches(0.25)
            ctf.margin_right = Inches(0.25)
            ctf.margin_top = Inches(0.15)
            for i, q in enumerate(s["mastery_questions"]):
                p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
                p.text = f"{i+1}.  {q}"
                p.font.size = Pt(13.5)
                p.font.color.rgb = NAVY
                p.space_after = Pt(7)
                p.alignment = PP_ALIGN.LEFT

        else:
            body_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), SLIDE_W - Inches(1.4), Inches(3.0))
            tf = body_box.text_frame
            tf.word_wrap = True
            for i, bullet in enumerate(s["concept"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = "•  " + bullet
                p.font.size = Pt(18)
                p.font.color.rgb = TEXT_DARK
                p.space_after = Pt(10)

            ex_top = Inches(1.4) + Inches(0.56) * len(s["concept"]) + Inches(0.35)
            ex_top = min(ex_top, Inches(5.1))
            ex_box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), ex_top, SLIDE_W - Inches(1.4), Inches(1.7)
            )
            ex_box.fill.solid()
            ex_box.fill.fore_color.rgb = AMBER_BG
            ex_box.line.color.rgb = AMBER
            ex_box.line.width = Pt(1.25)
            ex_box.shadow.inherit = False
            etf = ex_box.text_frame
            etf.word_wrap = True
            etf.vertical_anchor = MSO_ANCHOR.MIDDLE
            etf.margin_left = Inches(0.25)
            etf.margin_right = Inches(0.25)
            p = etf.paragraphs[0]
            run = p.add_run()
            run.text = "✏️ Exercise: " + s["exercise"]
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = AMBER

        add_slide_number(slide, idx, total)
        set_notes(slide, s["notes"])

    out_path = "Chapter3_Tensors_Practice_Presentation.pptx"
    prs.save(out_path)
    print(f"Saved {out_path} with {total} slides.")


if __name__ == "__main__":
    build()
