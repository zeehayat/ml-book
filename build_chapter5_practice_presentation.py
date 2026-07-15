"""
Build a "concept + exercise" PowerPoint presentation for Chapter 5
(Convex Optimization).

Run:
    python3 build_chapter5_practice_presentation.py

Produces:
    Chapter5_ConvexOpt_Practice_Presentation.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x1B, 0x26, 0x4F)
AMBER = RGBColor(0xB8, 0x6A, 0x00)
AMBER_BG = RGBColor(0xFD, 0xF2, 0xE2)
LIGHT_BG = RGBColor(0xF7, 0xF9, 0xFC)
TEXT_DARK = RGBColor(0x22, 0x22, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MASTERY_BG = RGBColor(0xEA, 0xF6, 0xF4)
TEAL = RGBColor(0x0D, 0x94, 0x88)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

SLIDES = [
    {
        "title": "Chapter 5: Convex Optimization",
        "subtitle": "One concept, one exercise per slide — practice as you go",
        "concept": [],
        "exercise": None,
        "notes": "Welcome slide. This chapter explains exactly when gradient descent is GUARANTEED to find the best possible answer, and how fast it gets there.",
    },
    {
        "title": "Concept: Convex Sets",
        "concept": [
            "A set is convex if the straight line between any two of its points stays inside the set",
            "Formally: for u, v in C and t in [0,1], the point t·u + (1-t)·v is also in C",
            "A disk is convex; a crescent (moon) shape is not",
        ],
        "exercise": "Is the set of all points with x ≥ 0 (a half-plane) convex? Is a donut shape (an annulus) convex?",
        "notes": "Answer: the half-plane IS convex — any line segment between two points with x≥0 stays in x≥0. The donut shape is NOT convex — a line segment between two points on opposite sides of the hole passes through the (excluded) hole.",
    },
    {
        "title": "Concept: Convex Functions (Jensen's Inequality)",
        "concept": [
            "f is convex if f(t·u + (1-t)·v) ≤ t·f(u) + (1-t)·f(v) for all u, v and t in [0,1]",
            "In words: the function's value on the line segment is never above the straight line connecting the endpoints",
            "Equivalent second-order test: the Hessian (second derivative matrix) is positive semidefinite everywhere",
        ],
        "exercise": "Using Jensen's inequality with u=0, v=4, t=0.5, check whether f(x)=x² is convex at this pair of points.",
        "notes": "Answer: f(0.5*0+0.5*4) = f(2) = 4. The right side: 0.5*f(0)+0.5*f(4) = 0.5*0+0.5*16 = 8. Since 4 ≤ 8, Jensen's inequality holds here — consistent with x² being convex everywhere.",
    },
    {
        "title": "Concept: Local Minima Are Global Minima (For Convex Functions)",
        "concept": [
            "The single most important theorem in this chapter",
            "If f is convex, ANY local minimum is automatically the GLOBAL minimum",
            "This is why gradient descent on a convex loss is trustworthy, not just hopeful",
            "For non-convex functions (most neural networks), this guarantee simply does not exist",
        ],
        "exercise": "Explain, in one sentence, why this theorem is the mathematical reason linear regression's normal equations and gradient descent always agree on the same answer.",
        "notes": "Answer: because the mean-squared-error loss is convex, its unique stationary point (found by either the closed-form normal equations or by gradient descent converging to a local minimum) must be the global minimum — there is no other valley for either method to land in.",
    },
    {
        "title": "Concept: Strict and Strong Convexity",
        "concept": [
            "Strict convexity: Jensen's inequality is STRICT — guarantees at most one global minimizer",
            "Strong convexity (μ > 0): quantifies HOW sharply the function curves upward from its minimum",
            "L2 regularization (weight_decay) adds λI to the Hessian, guaranteeing strong convexity",
            "This is why ridge regression always has a unique solution, even with collinear features",
        ],
        "exercise": "Plain (unregularized) linear regression on perfectly collinear features has infinitely many equally-good solutions. Why does adding an L2 penalty fix this?",
        "notes": "Answer: the unregularized Hessian X^T X can have a zero eigenvalue under perfect collinearity, so the objective is only weakly (not strictly) convex along that direction — many solutions tie. Adding lambda*I shifts every eigenvalue up by lambda, making the Hessian strictly positive definite and guaranteeing exactly one minimizer.",
    },
    {
        "title": "Concept: Gradient Descent's Safe Step Size",
        "concept": [
            "L-smoothness: the gradient cannot change faster than rate L (a Lipschitz bound)",
            "The descent lemma proves: any step size η < 2/L guarantees the loss decreases every step",
            "Too large a learning rate (η > 2/L) can cause the loss to increase, even on a convex function",
        ],
        "exercise": "A loss function has smoothness constant L=4. Which of these learning rates is guaranteed safe: η=0.3, η=0.5, η=0.6?",
        "notes": "Answer: the safe bound is η < 2/L = 2/4 = 0.5. So η=0.3 is safe; η=0.5 is exactly at the boundary (not strictly less than); η=0.6 exceeds the bound and is not guaranteed to decrease the loss every step.",
    },
    {
        "title": "Concept: The Condition Number and Convergence Speed",
        "concept": [
            "κ = λ_max / λ_min — the ratio of the Hessian's largest to smallest eigenvalue",
            "κ ≈ 1: a round bowl — gradient descent converges in very few steps",
            "κ ≫ 1: an elongated bowl — gradient descent zig-zags and converges slowly",
            "Feature scaling reshapes an elongated bowl back toward round, directly speeding up training",
        ],
        "exercise": "Two features are on wildly different scales: one ranges 0-1, the other ranges 0-100,000. Predict, qualitatively, whether training will converge quickly or slowly before any fix is applied — and name the one-line fix.",
        "notes": "Answer: slowly — the huge scale mismatch produces an elongated, poorly conditioned loss surface (large kappa), causing gradient descent to zig-zag. The fix is feature standardization (scaling both features to comparable ranges) before training, which reduces the condition number.",
    },
    {
        "title": "Concept: Newton's Method — Using Curvature",
        "concept": [
            "Gradient descent only uses the SLOPE (first derivative) at the current point",
            "Newton's method also uses the CURVATURE (the Hessian) to jump straight to the minimum",
            "On a true quadratic function, Newton's method converges in exactly ONE step, any κ",
            "Cost: forming and inverting the Hessian is O(D³) — prohibitive for large D",
        ],
        "exercise": "Why does modern large-scale machine learning almost universally use gradient descent (or its variants) instead of Newton's method, despite Newton's method needing far fewer steps?",
        "notes": "Answer: Newton's method's O(D^3) per-step cost to invert a D×D Hessian is completely impractical when D (the parameter count) is in the millions or billions — gradient descent's O(D) per-step cost, even across many more steps, remains vastly cheaper in total.",
    },
    {
        "title": "Concept: Historical Origins — Cauchy to Robbins-Monro",
        "concept": [
            "1847: Cauchy introduces gradient (steepest) descent for solving astronomical equations",
            "1805/1795: Legendre and Gauss solve least squares in closed form — no iteration needed",
            "1951: Robbins and Monro formalize stochastic approximation — the ancestor of mini-batch SGD",
        ],
        "exercise": "Why was gradient descent not necessary at all for solving the earliest least-squares problems, even though it was invented decades later?",
        "notes": "Answer: least squares (linear regression) has a closed-form solution via the normal equations — no iterative method is needed for a purely quadratic objective. Gradient descent became essential only once models grew nonlinear or too large for a closed-form matrix inversion to be practical.",
    },
    {
        "title": "Mastery Check: Are You Ready to Move On?",
        "concept": [
            "If you can answer every question below confidently, you have mastered Chapter 5.",
        ],
        "exercise": None,
        "mastery_questions": [
            "Define a convex set and a convex function, both geometrically and via Jensen's inequality.",
            "State and prove (in outline) why every local minimum of a convex function is global.",
            "Explain the difference between strict convexity and strong convexity.",
            "Explain why adding L2 regularization guarantees a unique minimizer even under collinearity.",
            "Derive the safe learning rate bound η < 2/L from the descent lemma.",
            "Define the condition number and explain its geometric meaning and effect on convergence speed.",
            "Explain why feature scaling speeds up gradient descent in terms of the condition number.",
            "Derive why Newton's method converges in exactly one step on a quadratic function.",
            "Explain why Newton's method is impractical at the scale of modern neural networks.",
            "Implement gradient descent from scratch and empirically verify the η < 2/L convergence boundary.",
        ],
        "notes": "This is a self-check slide. If any question causes hesitation, return to that concept's slide and its exercise before moving to Chapter 6.",
    },
]


def add_slide_number(slide, n, total):
    box = slide.shapes.add_textbox(SLIDE_W - Inches(1.3), SLIDE_H - Inches(0.5), Inches(1.0), Inches(0.4))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{n} / {total}"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def set_notes(slide, text):
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = text


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]
    total = len(SLIDES)

    for idx, s in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG

        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.15))
        header.fill.solid()
        header.fill.fore_color.rgb = NAVY
        header.line.fill.background()
        header.shadow.inherit = False
        title_box = header.text_frame
        title_box.word_wrap = True
        title_box.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = title_box.paragraphs[0]
        p.text = s["title"]
        p.font.size = Pt(28) if idx > 1 else Pt(32)
        p.font.bold = True
        p.font.color.rgb = WHITE
        title_box.margin_left = Inches(0.5)

        if idx == 1:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), SLIDE_W - Inches(1.6), Inches(1.5))
            tf = sub_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = s["subtitle"]
            run.font.size = Pt(22)
            run.font.color.rgb = TEAL
            run.font.italic = True
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.7), Inches(3.0), Pt(4))
            accent.fill.solid()
            accent.fill.fore_color.rgb = TEAL
            accent.line.fill.background()

        elif "mastery_questions" in s:
            body_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.35), SLIDE_W - Inches(1.2), Inches(0.5))
            tf = body_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = s["concept"][0]
            run.font.size = Pt(15)
            run.font.italic = True
            run.font.color.rgb = TEAL

            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.85), SLIDE_W - Inches(1.2), Inches(5.35)
            )
            box.fill.solid()
            box.fill.fore_color.rgb = MASTERY_BG
            box.line.color.rgb = TEAL
            box.line.width = Pt(1.25)
            box.shadow.inherit = False
            ctf = box.text_frame
            ctf.word_wrap = True
            ctf.margin_left = Inches(0.25)
            ctf.margin_right = Inches(0.25)
            ctf.margin_top = Inches(0.15)
            for i, q in enumerate(s["mastery_questions"]):
                p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
                p.text = f"{i+1}.  {q}"
                p.font.size = Pt(14.5)
                p.font.color.rgb = NAVY
                p.space_after = Pt(9)
                p.alignment = PP_ALIGN.LEFT

        else:
            body_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), SLIDE_W - Inches(1.4), Inches(3.0))
            tf = body_box.text_frame
            tf.word_wrap = True
            for i, bullet in enumerate(s["concept"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = "•  " + bullet
                p.font.size = Pt(18)
                p.font.color.rgb = TEXT_DARK
                p.space_after = Pt(10)

            ex_top = Inches(1.4) + Inches(0.56) * len(s["concept"]) + Inches(0.35)
            ex_top = min(ex_top, Inches(5.1))
            ex_box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), ex_top, SLIDE_W - Inches(1.4), Inches(1.7)
            )
            ex_box.fill.solid()
            ex_box.fill.fore_color.rgb = AMBER_BG
            ex_box.line.color.rgb = AMBER
            ex_box.line.width = Pt(1.25)
            ex_box.shadow.inherit = False
            etf = ex_box.text_frame
            etf.word_wrap = True
            etf.vertical_anchor = MSO_ANCHOR.MIDDLE
            etf.margin_left = Inches(0.25)
            etf.margin_right = Inches(0.25)
            p = etf.paragraphs[0]
            run = p.add_run()
            run.text = "✏️ Exercise: " + s["exercise"]
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = AMBER

        add_slide_number(slide, idx, total)
        set_notes(slide, s["notes"])

    out_path = "Chapter5_ConvexOpt_Practice_Presentation.pptx"
    prs.save(out_path)
    print(f"Saved {out_path} with {total} slides.")


if __name__ == "__main__":
    build()
